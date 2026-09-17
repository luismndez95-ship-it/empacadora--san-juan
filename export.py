"""Exportacion de la nomina semanal a Excel y PDF.

Separado de main.py para no inflar el archivo de rutas -- este modulo
solo sabe convertir un payroll_summary (de payroll.py) en bytes
descargables.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

import payroll

HEADER_FILL = PatternFill(start_color="2F6E42", end_color="2F6E42", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_excel(summary: dict, business_name: str, week_label: str, currency_symbol: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nomina Semanal"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"{business_name} - Nomina {week_label}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="left")

    headers = ["Empleado", "Horas regulares", "Horas extra", "Total horas", "Pago regular", "Pago extra", "Pago total"]
    ws.append([])
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in summary["rows"]:
        ws.append([
            row["full_name"],
            round(row["regular_hours"], 2),
            round(row["overtime_hours"], 2),
            round(row["total_hours"], 2),
            round(row["regular_pay"], 2),
            round(row["overtime_pay"], 2),
            round(row["total_pay"], 2),
        ])

    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=round(summary["grand_total_hours"], 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=round(summary["grand_total_pay"], 2)).font = Font(bold=True)

    for col_letter, width in zip("ABCDEFG", [26, 16, 14, 14, 16, 14, 16]):
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(summary: dict, business_name: str, week_label: str, currency_symbol: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Title"], textColor=colors.HexColor("#2F6E42"))

    elements = [
        Paragraph(business_name, title_style),
        Paragraph(f"Reporte de nomina semanal &mdash; {week_label}", styles["Normal"]),
        Spacer(1, 0.6 * cm),
    ]

    data = [["Empleado", "Hrs reg.", "Hrs extra", "Total hrs", "Pago reg.", "Pago extra", "Pago total"]]
    for row in summary["rows"]:
        data.append([
            row["full_name"],
            f"{row['regular_hours']:.2f}",
            f"{row['overtime_hours']:.2f}",
            f"{row['total_hours']:.2f}",
            payroll.format_crc(row["regular_pay"], currency_symbol),
            payroll.format_crc(row["overtime_pay"], currency_symbol),
            payroll.format_crc(row["total_pay"], currency_symbol),
        ])
    data.append([
        "TOTAL", "", "", f"{summary['grand_total_hours']:.2f}", "", "",
        payroll.format_crc(summary["grand_total_pay"], currency_symbol),
    ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F6E42")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5C0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F4F8F2")]),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)

    doc.build(elements)
    return buf.getvalue()
